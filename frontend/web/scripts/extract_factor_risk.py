"""
One-time extractor: Axioma weekly Time Series Risk Report (xlsx) -> src/data/factorRisk.json
Keeps the web app static (no xlsx parsing in the browser), matching the stfHistory.json pattern.

Run from frontend/web:  py scripts/extract_factor_risk.py
"""
import json
import os
import openpyxl

SRC = os.path.join(
    os.path.dirname(__file__), "..", "..", "..", "data", "2026-05-31",
    "weekly_US_EQ_Time Series Risk Report - Sortable_2026-03-31_2026-06-05.xlsx",
)
OUT = os.path.join(os.path.dirname(__file__), "..", "src", "data", "factorRisk.json")


def r(x, n=4):
    return None if not isinstance(x, (int, float)) else round(float(x), n)


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)

    def rows(name):
        ws = wb[name]
        data = list(ws.iter_rows(values_only=True))
        hdr = data[0]
        body = [row for row in data[1:] if len(row) > 2 and row[2]]  # has Period Description
        idx = {h: i for i, h in enumerate(hdr)}
        return hdr, body, idx

    _, rd, ri = rows("Risk Details")
    _, ae, ai = rows("Active Factor Exposure")
    _, fv, fvi = rows("Active Factor Risk (% of Var)")
    _, rt, rti = rows("Return Details")
    _, pd, pdi = rows("PortfolioDetails")

    def col(body, idx, name):
        return [r(row[idx[name]]) for row in body]

    def value(row, idx, name, default=0.0):
        i = idx.get(name)
        if i is None or i >= len(row):
            return default
        return row[i]

    dates = [row[2] for row in rd]

    # ---- per-period risk series ----
    series = []
    for k, row in enumerate(rd):
        series.append({
            "date": row[2],
            "predictedBeta": r(row[ri["Predicted Beta"]]),
            "historicalBeta": r(row[ri["Historical Beta"]]),
            "activeBeta": r(row[ri["Active Predicted Beta"]]),
            "totalRisk": r(row[ri["Total Predicted Risk"]]),
            "factorRisk": r(row[ri["Factor Risk"]]),
            "specificRisk": r(row[ri["Specific Risk"]]),
            "activeRisk": r(row[ri["Active Predicted Risk"]]),
            "activeFactorRisk": r(row[ri["Active Factor Risk"]]),
            "activeSpecificRisk": r(row[ri["Active Specific Risk"]]),
            "activeStyleRisk": r(row[ri["Active Style Factor Risk"]]),
            "activeIndustryRisk": r(row[ri["Active Industry Factor Risk"]]),
            "activeMarketRisk": r(value(row, ri, "Active Market Factor Risk")),
            "activeShare": r(row[ri["Active Share"]]),
        })

    # ---- return attribution (per-period + cumulative) ----
    def cum(name):
        p = 1.0
        for row in rt:
            v = row[rti[name]]
            if isinstance(v, (int, float)):
                p *= (1 + v)
        return round(p - 1, 4)

    def cum_sum(name):  # arithmetic sum for additive attribution lines
        i = rti.get(name)
        if i is None:
            return 0.0
        return round(sum(row[i] for row in rt if i < len(row) and isinstance(row[i], (int, float))), 4)

    returns = [{
        "date": row[2],
        "portfolio": r(row[rti["Period Return"]], 5),
        "benchmark": r(row[rti["Benchmark Period Return"]], 5),
        "active": r(row[rti["Active Period Return"]], 5),
    } for row in rt]

    attribution = {
        "portfolio": cum("Period Return"),
        "benchmark": cum("Benchmark Period Return"),
        "active": cum("Active Period Return"),
        "activeFactor": cum_sum("Active Period Factor Return"),
        "activeSpecific": cum_sum("Active Period Specific Return"),
        "activeStyle": cum_sum("Active Style Factor Returns"),
        "activeIndustry": cum_sum("Active Industry Factor Returns"),
        "activeMarket": cum_sum("Active Market Factor Returns"),
        "periods": len(rt),
        "from": rt[0][2],
        "to": rt[-1][2],
    }

    # ---- active style exposures (12 styles): current + start + drift series ----
    STYLE_COLS = [h for h in ai if isinstance(h, str) and h.startswith("Style Factors - ")]
    styleFactors = []
    for h in STYLE_COLS:
        vals = col(ae, ai, h)
        styleFactors.append({
            "name": h.replace("Style Factors - ", ""),
            "current": vals[-1],
            "start": vals[0],
            "series": vals,
        })
    styleFactors.sort(key=lambda f: -abs(f["current"] or 0))

    # ---- top active risk contributors (% of variance, last period) ----
    last = fv[-1]
    contributors = []
    for h, i in fvi.items():
        if not isinstance(h, str) or h in ("Period (Beginning)", "Period \n(Ending)", "Period Description"):
            continue
        v = last[i]
        if not isinstance(v, (int, float)):
            continue
        grp = next(
            (name for name in ("Style", "Country", "Industry", "Currency", "Local", "Market") if h.startswith(name)),
            "Other",
        )
        contributors.append({
            "name": h.split(" - ", 1)[1] if " - " in h else h,
            "group": grp,
            "pctVar": round(v, 4),
            "isCov": "Covariance" in h,
        })
    contributors.sort(key=lambda c: -(c["pctVar"] or 0))

    settings = {}
    for row in wb["Report Settings"].iter_rows(values_only=True):
        text = str(row[0] or "").strip() if row else ""
        if ":" in text:
            key, val = text.split(":", 1)
            settings[key.strip()] = val.strip()

    out = {
        "meta": {
            "portfolio": settings.get("Portfolio", ""),
            "benchmark": settings.get("Benchmark", ""),
            "riskModel": settings.get("Risk Model", ""),
            "currency": settings.get("Base Currency", ""),
            "from": dates[0],
            "to": dates[-1],
            "periods": len(dates),
            "holdings": pd[-1][pdi["Total # of\nAssets"]],
            "commonAssetsWeight": r(pd[-1][pdi["% Weight of Common Assets"]]),
        },
        "dates": dates,
        "series": series,
        "returns": returns,
        "attribution": attribution,
        "styleExposure": {"asOf": dates[-1], "factors": styleFactors},
        "contributors": contributors,
    }

    with open(OUT, "w") as fh:
        json.dump(out, fh, indent=2)
    print(f"wrote {OUT}  ({len(series)} periods, {len(styleFactors)} styles, {len(contributors)} contributors)")


if __name__ == "__main__":
    main()
